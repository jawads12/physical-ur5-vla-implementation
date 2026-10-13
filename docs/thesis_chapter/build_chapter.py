#!/usr/bin/env python3
"""Rebuild chapter figures, source register, HTML and editable Word document.

Uses Graphviz and system matplotlib. DOCX/HTML dependencies are installed
under this repository's .python_deps directory.
Does not import ROS or execute any robot-control code.
"""
from pathlib import Path
import csv
import hashlib
import json
import re
import subprocess
import sys

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
FIG = ROOT / 'figures'
FIG.mkdir(exist_ok=True)
REPO = PROJECT / 'ros2_ws/src/ur5-vla-dataset-recorder'
DATA = PROJECT / 'data/dataset'
MODELS = PROJECT / 'data/models'
OLD = PROJECT / 'dump/downloaded_models/drive-download-20260823T063551Z-1-001'
MODEL = PROJECT / 'dump/downloaded_models/drive-download-20260902T150120Z-1-001/ur5_cube_clean_73ep_v2_smolvla_a100_b96_16500_steps/pretrained_model'

STYLE = '''
graph [bgcolor="white", pad="0.25", nodesep="0.3", ranksep="0.45", dpi=180];
node [shape=box, style="rounded,filled", fillcolor="#edf3f9", color="#345878", fontname="DejaVu Sans", fontsize=12, margin="0.16,0.12"];
edge [color="#52677b", fontname="DejaVu Sans", fontsize=10, arrowsize=0.75];
'''
GRAPHS = {
'01_workflow': '''digraph workflow {
RANKSTYLE
rankdir=TB;
human [label="Human demonstration\nJoystick → MoveIt Servo → UR5"];
inputs [label="Two RGB views + aligned depth\nMeasured joints + gripper estimate\nFinal controller targets + language"];
raw [label="20 Hz raw episode recorder\nFrames, media, timing, metadata"];
gate [label="Structural + timing + behavior checks", fillcolor="#fff0d8"];
exclude [label="Document exclusions\nPreserve source records", fillcolor="#f7e7e7"];
convert [label="Versioned LeRobot dataset\nRGB video + 7D state/action + task"];
train [label="Training and checkpoint preservation\nWeights + configuration + processors"];
evaluate [label="Offline comparison → live shadow checks", fillcolor="#e7f3ea"];
future [label="Physical experiment protocol and trial records\nTask outcomes, timing, interventions", style="rounded,dashed", fillcolor="white"];
human -> inputs -> raw -> gate -> convert -> train -> evaluate -> future;
gate -> exclude [label="reject"];
{rank=same; gate; exclude;}
}''',
'03_control_architecture': '''digraph control {
RANKSTYLE
rankdir=TB;
sensors [label="ROS observation process\nBase/wrist RGB + joints + estimated gripper"];
model [label="Separate model process\nSmolVLA checkpoint → 50 × 7 actions"];
executor [label="Phased executor\nSelect ≤5 actions; validate; replace queue\nNominal publication: 20 Hz", fillcolor="#fff0d8"];
refs [label="Demonstration references\nApproach / carry / release locations"];
topic [label="/control_abstraction_layer\nSix absolute joints + binary gripper"];
engine [label="Control engine\nTarget checks + filtering\nVelocity/acceleration limits"];
arm [label="125 Hz arm commands\nForward position controller → UR5"];
io [label="Discrete gripper SetIO\nTwo sequential writes"];
ack [label="Acknowledged I/O state\nNot measured grasp success", fillcolor="#e7f3ea"];
sensors -> model [label="localhost request"];
model -> executor [label="prediction"];
refs -> executor;
executor -> topic -> engine;
engine -> arm;
engine -> io -> ack;
ack -> sensors [style=dashed, constraint=false];
{rank=same; arm; io;}
}''',
'04_phase_sequence': '''digraph phases {
RANKSTYLE
rankdir=TB;
approach [label="APPROACH\nGripper remains open"];
close [label="Closing request pending\nHold measured joint target", fillcolor="#fff0d8"];
carry [label="CARRY\nGripper remains closed"];
open [label="Opening request pending\nHold measured joint target", fillcolor="#fff0d8"];
released [label="RELEASED\nWait 2 seconds"];
complete [label="COMPLETE\nEnd of release sequence\nObject outcome requires observation", fillcolor="#e7f3ea"];
fault [label="FAULT\nClear queued actions\nStop executor publication", fillcolor="#f7e7e7"];
approach -> close [label="3 matching predictions\n+ demonstrated closing-region check"];
close -> carry [label="matching I/O acknowledgment\n+ ≥1 second settling"];
carry -> open [label="3 matching predictions\n+ demonstrated opening-region check"];
open -> released [label="matching I/O acknowledgment\n+ ≥1 second settling"];
released -> complete;
approach -> fault [style=dashed, label="prediction / runtime fault"];
carry -> fault [style=dashed];
{rank=same; carry; fault;}
}'''
}

for stem, code in GRAPHS.items():
    path = FIG / f'{stem}.dot'
    path.write_text(code.replace('RANKSTYLE', STYLE))
    for extension in ('svg', 'png', 'pdf'):
        subprocess.run(['dot', f'-T{extension}', str(path), '-o', str(FIG / f'{stem}.{extension}')], check=True)

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
reports = []
for path in sorted((MODELS / 'evaluation/53ep_checkpoint_comparison').glob('*.json')):
    d = json.loads(path.read_text())
    reports.append(dict(steps=int(path.stem), samples=d['samples'],
        mean_action_error=d['predicted_vs_demonstrated_action_error_rad']['mean'],
        mean_target_delta=d['predicted_target_delta_rad']['mean'],
        max_target_delta=d['predicted_target_delta_rad']['max'],
        failures=d['predicted_target_delta_rad']['failures'],
        limit=d['predicted_target_delta_rad']['limit'], passed=d['pass']))
with (ROOT / 'checkpoint_comparison.csv').open('w') as stream:
    writer = csv.DictWriter(stream, fieldnames=list(reports[0]))
    writer.writeheader(); writer.writerows(reports)
plt.rcParams.update({'font.size': 10, 'font.family': 'DejaVu Sans', 'axes.spines.top': False, 'axes.spines.right': False})
fig, (ax, bx) = plt.subplots(2, 1, figsize=(9, 6.8), sharex=True, gridspec_kw={'height_ratios': [2, 1]}, constrained_layout=True)
x = [r['steps']/1000 for r in reports]
ax.plot(x, [r['mean_action_error'] for r in reports], 'o-', color='#245782', label='Mean maximum-joint demonstration error E')
ax.plot(x, [r['mean_target_delta'] for r in reports], 's--', color='#c47122', markersize=4, label='Mean maximum-joint target displacement D')
ax.axhline(reports[0]['limit'], color='#954144', linestyle=':', label='Displacement limit: 0.01256637 rad')
ax.set(ylabel='Angle (rad)', title='Saved offline comparison: 53 samples per checkpoint')
ax.set_ylim(0, .10); ax.grid(alpha=.2); ax.legend(fontsize=9)
bx.bar(x, [r['failures'] for r in reports], width=3.2, color='#496b83')
bx.axhline(53, color='#954144', linestyle=':', linewidth=1)
bx.set(ylabel='Samples above limit', xlabel='Training steps (thousands)', ylim=(0, 58), xticks=x)
bx.grid(axis='y', alpha=.2)
for extension in ('svg', 'png', 'pdf'):
    fig.savefig(FIG / f'02_checkpoint_results.{extension}', dpi=240)
plt.close(fig)

groups = {
'E1': [REPO/'docs/PROJECT_CONTEXT.md'],
'E2': [REPO/'docs/CODEX_HANDOFF_2026-08-23.md'],
'E3': [REPO/'README.md', REPO/'config/recorder.yaml', REPO/'ur5_vla_dataset_recorder/core.py', REPO/'ur5_vla_dataset_recorder/recorder_node.py'],
'E4': [REPO/'docs/RAW_VALIDATION_RESULTS.md'],
'E5': [REPO/'config/ur5_cube_clean_20ep.json'],
'E6': sorted((DATA/'lerobot').glob('*/meta/info.json')) + sorted((DATA/'lerobot').glob('*/meta/ur5_vla_conversion.json')),
'E7': [REPO/'docs/PHYSICAL_UR5_TRAINING_RUNBOOK.md', REPO/'docs/LEROBOT_PIPELINE.md'],
'E8': [MODEL/'train_config.json', MODEL/'config.json'],
'E9': [REPO/'docs/LIVE_ONE_SHOT_EXECUTION.md', REPO/'docs/PHASED_DEPLOYMENT.md'],
'E10': sorted((PROJECT/'scripts').glob('*.py')),
'E11': [OLD/'offline_validation/report_100_frames.json', OLD/'offline_validation/actions_100_frames.csv'],
'E12': sorted((MODELS/'evaluation/53ep_checkpoint_comparison').glob('*.json')) + [REPO/'ur5_vla_dataset_recorder/offline_inference.py'],
'E13': [REPO/'ur5_vla_dataset_recorder/action_semantics.py', DATA/'lerobot/ur5_cube_clean_20ep_delta/delta_integrity_report.json'],
'E14': [REPO/'tools/generate_pi0_73ep_a100_colab.py']
}
register = {'inspection_date': '2026-09-30', 'evidence': {key: [dict(path=str(p), sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in paths] for key, paths in groups.items()}}
(ROOT/'evidence_register.json').write_text(json.dumps(register, indent=2)+'\n')

sys.path.insert(0, str(PROJECT/'.python_deps'))
import markdown
source = (ROOT/'experimental_methodology.md').read_text()
html = markdown.markdown(source, extensions=['tables', 'fenced_code'])
css = '''body{font:17px/1.65 Georgia,serif;max-width:980px;margin:50px auto;padding:0 28px;color:#202c35}
h1,h2,h3{font-family:Arial,sans-serif;line-height:1.25;color:#244d6b}h2{margin-top:2em}
img{display:block;max-width:100%;max-height:850px;margin:24px auto}table{border-collapse:collapse;width:100%;font:14px/1.45 Arial,sans-serif;margin:20px 0}
td,th{border:1px solid #bdcbd7;padding:9px;text-align:left}th{background:#edf3f9}pre{white-space:pre-wrap;background:#f2f5f7;padding:16px;font-size:15px}
@media print{body{max-width:none;margin:0;font-size:11pt}h2,h3{break-after:avoid}img{max-height:8in}tr{break-inside:avoid}pre{break-inside:avoid}@page{size:A4;margin:20mm}'''
(ROOT/'experimental_methodology.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>UR5 experimental platform chapter</title><style>'+css+'</style><body>'+html+'</body></html>')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
document = Document()
section = document.sections[0]
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.8)
normal = document.styles['Normal']
normal.font.name = 'Calibri'; normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.15
for name in ('Heading 1','Heading 2','Heading 3'):
    document.styles[name].font.color.rgb = RGBColor.from_string('244D6B')

def inline(paragraph, value):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]*`)', value):
        if not part: continue
        if part.startswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = paragraph.add_run(part[1:-1]); run.font.name = 'Consolas'; run.font.size = Pt(9)
        else:
            paragraph.add_run(part)

lines = source.splitlines(); i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line.startswith('#'):
        level = len(line)-len(line.lstrip('#'))
        document.add_heading(line[level:].strip(), level=0 if level == 1 else level-1)
    elif line.startswith('!['):
        match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        pic = document.add_picture(str(ROOT/match.group(2)))
        ratio = pic.width/pic.height
        pic.width = Inches(min(6.6, 6.2*ratio)); pic.height = int(pic.width/ratio)
        document.paragraphs[-1].alignment = 1
        document.paragraphs[-1].paragraph_format.keep_with_next = True
    elif line.startswith('```'):
        block=[]; i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            block.append(lines[i]); i += 1
        p=document.add_paragraph(); run=p.add_run('\n'.join(block)); run.font.name='DejaVu Sans Mono'; run.font.size=Pt(9)
    elif line.startswith('|'):
        rows=[]
        while i < len(lines) and lines[i].startswith('|'):
            cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', c) for c in cells): rows.append(cells)
            i+=1
        table=document.add_table(rows=1, cols=len(rows[0])); table.style='Light Shading Accent 1'
        for col,text in enumerate(rows[0]): inline(table.rows[0].cells[col].paragraphs[0],text)
        props=table.rows[0]._tr.get_or_add_trPr(); props.append(OxmlElement('w:tblHeader'))
        for row in rows[1:]:
            cells=table.add_row().cells
            for col,text in enumerate(row): inline(cells[col].paragraphs[0],text)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs: run.font.size=Pt(9)
        document.add_paragraph(); continue
    elif line.startswith('*Figure'):
        p=document.add_paragraph(); p.add_run(line.strip('*')).italic=True
    else:
        p=document.add_paragraph(); inline(p,line)
    i+=1
footer=section.footer.paragraphs[0]; footer.alignment=2
footer.add_run('Experimental platform draft • ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
document.core_properties.title='Development and Validation of the UR5 Experimental Platform'
document.core_properties.subject='Thesis chapter draft with project evidence and future experiment structure'
document.core_properties.author=''
document.save(ROOT/'experimental_methodology.docx')
print(json.dumps({'words':len(source.split()), 'checkpoint_reports':len(reports), 'sources':sum(map(len,groups.values())), 'output_dir':str(ROOT)},indent=2))
